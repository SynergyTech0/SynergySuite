"""Write a rich .docx with python-docx, the way a Word user's file looks:
built-in and custom styles, direct formatting, multilevel lists, a table with
merges and shading, a picture, a link, a page break, two sections (portrait
then landscape with two columns), headers and footers with page numbers,
a different first page, and a comment.

    python3 make_python_docx.py out.docx
"""
import struct
import sys
import zlib

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def png(w, h, rgb):
    raw = b"".join(b"\x00" + bytes(rgb) * w for _ in range(h))
    def chunk(t, d):
        return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")


def field(paragraph, instr):
    run = paragraph.add_run()
    for kind, text in (("begin", None), ("instr", instr), ("separate", None), ("text", "1"), ("end", None)):
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
            run._r.append(el)
        elif kind == "text":
            t = OxmlElement("w:t")
            t.text = text
            run._r.append(t)
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            run._r.append(el)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def hyperlink(paragraph, url, text):
    rid = paragraph.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rs = OxmlElement("w:rStyle")
    rs.set(qn("w:val"), "Hyperlink")
    rpr.append(rs)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    h.append(r)
    paragraph._p.append(h)


def main(out):
    doc = Document()
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Cambria"
    normal.font.size = Pt(11)

    custom = styles.add_style("Callout Box", WD_STYLE_TYPE.PARAGRAPH)
    custom.base_style = normal
    custom.font.name = "Georgia"
    custom.font.size = Pt(13)
    custom.font.bold = True
    custom.font.color.rgb = RGBColor(0x33, 0x66, 0x99)
    custom.paragraph_format.space_after = Pt(18)
    custom.paragraph_format.left_indent = Pt(24)
    custom.quick_style = True
    cstyle = styles.add_style("Red Note", WD_STYLE_TYPE.CHARACTER)
    cstyle.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    cstyle.font.italic = True

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    sec.left_margin = sec.right_margin = Inches(0.75)
    sec.different_first_page_header_footer = True
    sec.first_page_header.paragraphs[0].text = "Cover header"
    sec.header.paragraphs[0].text = "Running header"
    sec.header.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run("Page ")
    field(fp, " PAGE ")
    fp.add_run(" of ")
    field(fp, " NUMPAGES ")

    doc.add_paragraph("Fixture Report", style="Title")
    doc.add_heading("Introduction", level=1)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.first_line_indent = Pt(18)
    p.paragraph_format.line_spacing = 1.5
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run("Bold").bold = True
    p.add_run(" and ")
    r = p.add_run("blue Arial 14")
    r.font.name = "Arial"
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(0x00, 0x00, 0xFF)
    p.add_run(" and ")
    r = p.add_run("highlighted")
    r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    p.add_run(" and ")
    p.add_run("a red note", style="Red Note")
    p.add_run(" and ")
    hyperlink(p, "https://example.org/docs", "a link")
    r = p.add_run(" commented text")
    doc.add_comment(r, text="Please check this sentence.", author="Grace Hopper", initials="GH")

    doc.add_heading("Lists", level=2)
    doc.add_paragraph("First step", style="List Number")
    doc.add_paragraph("Second step", style="List Number")
    doc.add_paragraph("Detail under second", style="List Number 2")
    doc.add_paragraph("A bullet", style="List Bullet")
    doc.add_paragraph("A nested bullet", style="List Bullet 2")

    doc.add_paragraph("Callout paragraph in a custom style", style="Callout Box")

    table = doc.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    for i, row in enumerate(table.rows):
        for j, c in enumerate(row.cells):
            c.text = f"r{i}c{j}"
    a = table.cell(0, 0).merge(table.cell(0, 1))
    a.text = "merged across"
    b = table.cell(1, 2).merge(table.cell(2, 2))
    b.text = "merged down"
    shade(table.cell(1, 0), "D9E2F3")

    import io
    doc.add_picture(io.BytesIO(png(40, 20, (200, 30, 30))), width=Inches(2))

    br = doc.add_paragraph().add_run()
    br.add_break(WD_BREAK.PAGE)
    doc.add_heading("Second page", level=3)
    doc.add_paragraph("Text on the second page.")

    new = doc.add_section(WD_SECTION.NEW_PAGE)
    new.orientation = WD_ORIENT.LANDSCAPE
    new.page_width, new.page_height = Inches(11.69), Inches(8.27)
    cols = new._sectPr.xpath("./w:cols")
    col = cols[0] if cols else OxmlElement("w:cols")
    col.set(qn("w:num"), "2")
    col.set(qn("w:space"), "720")
    if not cols:
        new._sectPr.append(col)
    doc.add_heading("Landscape section", level=1)
    for i in range(6):
        doc.add_paragraph(f"Column text paragraph {i + 1}. " * 6)

    doc.core_properties.title = "Fixture Report"
    doc.save(out)


if __name__ == "__main__":
    main(sys.argv[1])
