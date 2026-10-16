"""Writes a small deck with python-pptx (default PowerPoint template) to argv[1].

Used by tests/pptx-import.test.ts when python-pptx is installed, so the
importer is exercised against the placeholder/master structure PowerPoint
itself produces. Nothing binary is committed.
"""
import base64
import io
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAIAAAACCAIAAAD91JpzAAAAEElEQVR4nGP4z8AARAwQCgAf7gP9i18U1AAAAABJRU5ErkJggg=="
)

prs = Presentation()

s = prs.slides.add_slide(prs.slide_layouts[0])
s.shapes.title.text = "Quarterly Review"
s.placeholders[1].text = "Prepared by the team"

s = prs.slides.add_slide(prs.slide_layouts[1])
s.shapes.title.text = "Highlights"
tf = s.placeholders[1].text_frame
tf.text = "Revenue up"
p = tf.add_paragraph()
p.text = "Mostly in EMEA"
p.level = 1
p = tf.add_paragraph()
r = p.add_run()
r.text = "Costs flat"
r.font.bold = True
r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
s.notes_slide.notes_text_frame.text = "Mention the hiring plan"

s = prs.slides.add_slide(prs.slide_layouts[6])
tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
p = tb.text_frame.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "Big italic"
r.font.size = Pt(40)
r.font.italic = True
r.font.underline = True
shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5), Inches(1), Inches(2), Inches(1))
shp.fill.solid()
shp.fill.fore_color.rgb = RGBColor(0x00, 0x80, 0x00)
shp.text = "Go"
s.shapes.add_picture(io.BytesIO(PNG), Inches(1), Inches(3), Inches(1), Inches(1))
tbl = s.shapes.add_table(2, 2, Inches(4), Inches(3), Inches(4), Inches(1)).table
tbl.cell(0, 0).text = "Q"
tbl.cell(0, 1).text = "Total"
tbl.cell(1, 0).text = "Q1"
tbl.cell(1, 1).text = "42"

prs.save(sys.argv[1])
